import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


def relevant(majors, requirements):
    return (
        "电子/通信/自动化" in majors
        or "电气/能源/动力类" in majors
        or re.search(r"不限专业|专业不限|理工科", f"{majors} {requirements}")
    )


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("events")
    args = parser.parse_args()

    sheet = load_workbook(args.workbook, read_only=True, data_only=True).active
    company = ""
    companies = []
    jobs = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        company = row[1] or company
        title = str(row[22] or "").strip()
        requirements = str(row[23] or "")
        majors = str(row[24] or "").strip()
        if not title or not relevant(majors, requirements):
            continue
        if company not in companies:
            companies.append(company)
        jobs.append({
            "title": title,
            "company": company,
            "salary": f"{str(row[27]).strip()}元/月" if row[27] else None,
            "education": str(row[25]).strip() if row[25] else None,
            "majors": [item.strip() for item in majors.split(",") if item.strip()],
            "location": str(row[29]).strip() if row[29] else None,
            "headcount": int(row[26]) if str(row[26] or "").strip().isdigit() else None,
        })

    event = {
        "id": "xiamen-tongan-fjut-20260923",
        "title": "职选厦一站·未来无限量——2026年秋季厦门（同安）校园招聘会（福建理工大学站）",
        "eventType": "双选会",
        "startAt": "2026-09-23T14:30:00+08:00",
        "endAt": "2026-09-23T17:00:00+08:00",
        "mode": "线下",
        "venue": "福建理工大学旗山校区北区风雨篮球场",
        "companies": companies,
        "jobs": jobs,
        "majorTags": ["电气/能源/动力类", "电子/通信/自动化", "理工科", "不限专业"],
        "applicationUrl": "https://app.xmrc.com.cn/JobFair/MeetingInfo/3bc3642321d4403b82d41162b37ec29a",
        "sourceUrl": "https://mp.weixin.qq.com/s/swv7HJ0EkEFJm9FqgBTSQQ",
        "sourceName": "厦门市人才服务中心",
        "reviewStatus": "approved",
        "reviewedAt": "2026-09-22T16:30:00+08:00",
        "updatedAt": "2026-09-22T16:30:00+08:00",
    }

    path = Path(args.events)
    events = json.loads(path.read_text(encoding="utf-8"))
    events = [item for item in events if item["id"] not in {event["id"], "xiamen-hnu-20260922"}]
    events.append(event)
    path.write_text(json.dumps(events, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"jobs": len(jobs), "companies": len(companies)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
